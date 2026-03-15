import os
import re
import json
import uuid
from datetime import datetime
from html import escape
from collections import Counter
from fastapi import FastAPI, UploadFile, File, Form, Request, Query
from fastapi.staticfiles import StaticFiles
from fastapi.responses import HTMLResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from sentence_transformers import SentenceTransformer
from sklearn.metrics.pairwise import cosine_similarity
from search_utils import search_internet_for_sentence_async
import numpy as np
import warnings

import fitz # PyMuPDF
import asyncio
import aiohttp
from docx import Document
import io
from threading import Lock

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.graphics.shapes import Drawing
    from reportlab.graphics.charts.barcharts import VerticalBarChart
    REPORTLAB_AVAILABLE = True
except Exception:
    REPORTLAB_AVAILABLE = False

# Suppress annoying warnings
warnings.filterwarnings("ignore")

app = FastAPI(title="Advanced AI Plagiarism Checker")
MAX_WEB_QUERIES = 40
REPORT_CACHE = {}
MAX_REPORT_CACHE_ITEMS = 50
ASSET_VERSION = datetime.utcnow().strftime("%Y%m%d%H%M%S")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

model = None
internal_embeddings = None
ai_reference_embeddings = None
semantic_reference_embeddings = None
_model_init_lock = Lock()

INTERNAL_DB = [
    "Artificial intelligence is transforming industries and modern businesses.",
    "Machine learning allows computers to learn from data automatically.",
    "Deep learning is a subset of machine learning using neural networks.",
    "Machine learning allows computers to learn from data without explicit programming."
]

AI_REFERENCE_DB = [
    "In conclusion, this approach offers a scalable and efficient solution for modern organizations.",
    "Furthermore, it is important to note that implementing this strategy can significantly improve productivity.",
    "Overall, this project demonstrates the potential of artificial intelligence in real-world applications.",
    "The system can be integrated with existing tools to provide seamless automation and better decision-making.",
    "This solution is user-friendly, cost-effective, and capable of handling large volumes of data.",
    "As a result, businesses can reduce manual effort and focus on higher-value tasks.",
    "The proposed model leverages machine learning to optimize performance and deliver accurate predictions.",
    "In summary, this framework provides a practical pathway to digital transformation.",
    "This project can be extended in the future by adding advanced analytics and personalized recommendations.",
    "The architecture is robust, modular, and suitable for deployment in cloud environments.",
    "Here are some project ideas you can build using this technology.",
    "Why it is good: it is practical, easy to implement, and demonstrates end-to-end AI capabilities.",
]
SEMANTIC_REFERENCE_DB = INTERNAL_DB + AI_REFERENCE_DB


def ensure_model_resources_loaded():
    global model, internal_embeddings, ai_reference_embeddings, semantic_reference_embeddings

    if model is not None:
        return

    with _model_init_lock:
        if model is not None:
            return

        print("Loading SentenceTransformer model...")
        loaded_model = SentenceTransformer('all-MiniLM-L6-v2')
        loaded_internal_embeddings = loaded_model.encode(INTERNAL_DB)
        loaded_ai_reference_embeddings = loaded_model.encode(AI_REFERENCE_DB)
        loaded_semantic_embeddings = loaded_model.encode(SEMANTIC_REFERENCE_DB)

        model = loaded_model
        internal_embeddings = loaded_internal_embeddings
        ai_reference_embeddings = loaded_ai_reference_embeddings
        semantic_reference_embeddings = loaded_semantic_embeddings
        print("Model loaded.")

app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/", response_class=HTMLResponse)
async def read_index():
    with open("static/index.html", "r", encoding="utf-8") as f:
        html = f.read()

    # Cache-bust static assets so browser always loads the latest frontend code.
    html = html.replace(
        '/static/script.js',
        f'/static/script.js?v={ASSET_VERSION}'
    )
    html = html.replace(
        '/static/styles.css',
        f'/static/styles.css?v={ASSET_VERSION}'
    )

    return HTMLResponse(
        content=html,
        headers={
            "Cache-Control": "no-store, no-cache, must-revalidate, max-age=0",
            "Pragma": "no-cache",
            "Expires": "0",
        },
    )

def extract_text_from_pdf(file_bytes):
    text = ""
    try:
        pdf_document = fitz.open(stream=file_bytes, filetype="pdf")
        for page_num in range(len(pdf_document)):
            page = pdf_document.load_page(page_num)
            text += page.get_text() + "\n"
        pdf_document.close()
    except Exception as e:
        print(f"Error reading PDF with PyMuPDF: {e}")
    return text

def extract_text_from_docx(file_bytes):
    doc = Document(io.BytesIO(file_bytes))
    return "\n".join([p.text for p in doc.paragraphs])

def split_into_sentences(text):
    # Remove weird PDF artifacts and extra linebreaks
    text = re.sub(r'[\r\n]+', ' ', text)
    text = re.sub(r'\s+', ' ', text)
    
    sentences = re.split(r'(?<=[.!?])\s+', text)
    processed = []
    
    # Batch them into slightly larger chunks so we don't query DDG 500 times
    current_chunk = ""
    for s in sentences:
        s = s.strip()
        if len(s) < 15: # Skip tiny fragments
            continue
            
        if len(current_chunk) + len(s) < 150: # Group sentences up to 150 chars
            current_chunk += s + " "
        else:
            if current_chunk:
                processed.append(current_chunk.strip())
            current_chunk = s + " "
            
    if current_chunk:
        processed.append(current_chunk.strip())
        
    # Cap maximum sentences to check at 100 to prevent loading times on huge PDFs from taking too long
    return processed[:100]

def analyze_ai_usage(sentences, source_text):
    """
    Lightweight heuristic AI-writing detector.
    Returns probabilistic signals, not a definitive classifier.
    """
    normalized_text = re.sub(r"\s+", " ", source_text).strip().lower()
    words = re.findall(r"\b[a-zA-Z']+\b", normalized_text)
    total_words = len(words)

    if total_words == 0 or not sentences:
        return {
            "ai_usage_score": 0.0,
            "ai_usage_level": "Unknown",
            "ai_likely_sentences": 0,
            "ai_likely_indices": [],
            "ai_sentence_details": {},
            "ai_analysis": {
                "confidence": "low",
                "reason": "Insufficient text for AI usage analysis.",
            },
        }

    sentence_lengths = [len(re.findall(r"\b\w+\b", s)) for s in sentences]
    avg_sentence_len = float(np.mean(sentence_lengths)) if sentence_lengths else 0.0
    sent_len_std = float(np.std(sentence_lengths)) if len(sentence_lengths) > 1 else 0.0
    burstiness = sent_len_std / (avg_sentence_len + 1e-6)

    unique_words = len(set(words))
    lexical_diversity = unique_words / total_words if total_words else 0.0

    transitions = [
        "moreover", "furthermore", "in addition", "additionally", "therefore",
        "in conclusion", "overall", "it is important to note", "on the other hand",
        "as a result", "consequently", "in summary", "to summarize", "thus"
    ]
    transition_hits = sum(normalized_text.count(phrase) for phrase in transitions)
    transition_density = transition_hits / max(len(sentences), 1)

    first_person = [
        " i ", " my ", " me ", " we ", " our ", " us "
    ]
    padded = f" {normalized_text} "
    first_person_hits = sum(padded.count(token) for token in first_person)
    first_person_rate = first_person_hits / max(total_words, 1)

    openings = []
    for sentence in sentences:
        tokens = re.findall(r"\b[a-zA-Z']+\b", sentence.lower())
        if len(tokens) >= 2:
            openings.append(" ".join(tokens[:2]))
    opening_counts = Counter(openings)
    repeated_openings = sum(c for c in opening_counts.values() if c > 1)
    repeated_opening_ratio = repeated_openings / max(len(sentences), 1)

    punctuation_density = source_text.count(",") / max(len(sentences), 1)
    question_exclaim_count = source_text.count("?") + source_text.count("!")
    list_item_markers = len(re.findall(r"\b\d+\.", source_text))

    sentence_embeddings = model.encode(sentences)
    ai_reference_sims = cosine_similarity(sentence_embeddings, ai_reference_embeddings)
    top_ai_similarities = np.max(ai_reference_sims, axis=1) if len(sentences) else np.array([0.0])
    avg_ai_reference_similarity = float(np.mean(top_ai_similarities))

    score = 0.0

    if burstiness < 0.35:
        score += 22
    elif burstiness < 0.5:
        score += 12

    if transition_density > 0.35:
        score += 22
    elif transition_density > 0.2:
        score += 12

    if first_person_rate < 0.012:
        score += 12

    if repeated_opening_ratio > 0.35:
        score += 28
    elif repeated_opening_ratio > 0.2:
        score += 20
    elif repeated_opening_ratio > 0.1:
        score += 12

    if punctuation_density > 1.2 and question_exclaim_count <= 1:
        score += 10

    if 0.35 <= lexical_diversity <= 0.6:
        score += 8
    elif lexical_diversity < 0.34:
        score += 8

    if list_item_markers >= 8:
        score += 16
    elif list_item_markers >= 4:
        score += 8

    if avg_ai_reference_similarity >= 0.66:
        score += 34
    elif avg_ai_reference_similarity >= 0.58:
        score += 24
    elif avg_ai_reference_similarity >= 0.5:
        score += 14

    if first_person_rate > 0.03:
        score -= 8

    ai_sentence_markers = [
        "in conclusion", "overall", "furthermore", "it is important to note",
        "in summary", "to summarize", "therefore"
    ]
    ai_likely_sentences = 0
    ai_likely_indices = []
    ai_sentence_details = {}
    for idx, sentence in enumerate(sentences):
        s = sentence.lower()
        marker_present = any(marker in s for marker in ai_sentence_markers)
        token_count = len(re.findall(r"\b[a-zA-Z']+\b", s))
        has_personal_voice = bool(re.search(r"\b(i|my|me|we|our|us)\b", s))
        sim_to_ai_reference = float(top_ai_similarities[idx]) if idx < len(top_ai_similarities) else 0.0
        reasons = []

        if marker_present:
            reasons.append("contains common AI-transition phrasing")
        if token_count >= 20 and not has_personal_voice:
            reasons.append("long formal sentence without personal voice")
        if sim_to_ai_reference >= 0.62:
            reasons.append(f"high AI-reference similarity ({round(sim_to_ai_reference * 100)}%)")

        if marker_present or (token_count >= 20 and not has_personal_voice) or sim_to_ai_reference >= 0.62:
            ai_likely_sentences += 1
            ai_likely_indices.append(idx)

        if sim_to_ai_reference >= 0.72:
            sentence_confidence = "high"
        elif sim_to_ai_reference >= 0.6:
            sentence_confidence = "medium"
        else:
            sentence_confidence = "low"

        ai_sentence_details[idx] = {
            "reasons": reasons,
            "confidence": sentence_confidence,
            "ai_reference_similarity": round(sim_to_ai_reference, 3),
        }

    ai_sentence_ratio = ai_likely_sentences / max(len(sentences), 1)
    if ai_sentence_ratio >= 0.35:
        score += 22
    elif ai_sentence_ratio >= 0.2:
        score += 14
    elif ai_sentence_ratio >= 0.1:
        score += 8

    score = max(0.0, min(100.0, score))

    if score >= 70:
        level = "High"
    elif score >= 40:
        level = "Moderate"
    else:
        level = "Low"

    if total_words < 120:
        confidence = "low"
        reason = "Short text lowers AI-detection reliability."
    elif total_words < 250:
        confidence = "medium"
        reason = "Moderate text length gives a directional estimate."
    else:
        confidence = "high"
        reason = "Sufficient text length for stronger pattern signals."

    return {
        "ai_usage_score": round(score, 1),
        "ai_usage_level": level,
        "ai_likely_sentences": ai_likely_sentences,
        "ai_likely_indices": ai_likely_indices,
        "ai_sentence_details": ai_sentence_details,
        "ai_analysis": {
            "confidence": confidence,
            "reason": reason,
            "metrics": {
                "average_sentence_length": round(avg_sentence_len, 2),
                "sentence_burstiness": round(burstiness, 3),
                "lexical_diversity": round(lexical_diversity, 3),
                "transition_density": round(transition_density, 3),
                "first_person_rate": round(first_person_rate, 3),
                "repeated_opening_ratio": round(repeated_opening_ratio, 3),
                "ai_reference_similarity": round(avg_ai_reference_similarity, 3),
                "ai_sentence_ratio": round(ai_sentence_ratio, 3),
                "list_item_markers": list_item_markers,
            },
        },
    }

def confidence_from_similarity(similarity):
    if similarity >= 0.75:
        return "high"
    if similarity >= 0.6:
        return "medium"
    return "low"

def build_report_html(data):
    results = data.get("results", []) if isinstance(data, dict) else []

    generated_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    plagiarism_score = round(float(data.get("plagiarism_score", 0) or 0))
    ai_usage_score = round(float(data.get("ai_usage_score", 0) or 0))
    ai_usage_level = escape(str(data.get("ai_usage_level", "Unknown")))
    plagiarized_sentences = int(data.get("plagiarized_sentences", 0) or 0)
    ai_likely_sentences = int(data.get("ai_likely_sentences", 0) or 0)
    total_sentences = int(data.get("total_sentences", 0) or 0)

    rendered_sentences = []
    flagged_details = []
    internal_groups = []
    internet_groups = []

    for res in results:
        sentence = escape(str(res.get("sentence", "")))
        is_plagiarized = bool(res.get("is_plagiarized"))
        is_ai_likely = bool(res.get("is_ai_likely"))

        css_class = "normal"
        if is_plagiarized and is_ai_likely:
            css_class = "both"
        elif is_plagiarized:
            css_class = "plagiarized"
        elif is_ai_likely:
            css_class = "ai"
        rendered_sentences.append(f'<span class="sentence {css_class}">{sentence}</span>')

        if is_plagiarized or is_ai_likely:
            reasons = res.get("flag_reasons", [])
            reason_text = "; ".join([escape(str(r)) for r in reasons]) if reasons else "Heuristic threshold triggered."
            sentence_confidence = escape(str(res.get("sentence_confidence", "low"))).capitalize()
            flagged_details.append(
                f"""
                <tr>
                    <td>{sentence}</td>
                    <td>{'Yes' if is_plagiarized else 'No'}</td>
                    <td>{'Yes' if is_ai_likely else 'No'}</td>
                    <td>{sentence_confidence}</td>
                    <td>{reason_text}</td>
                </tr>
                """
            )

        for src in res.get("sources", []):
            similarity = float(src.get("similarity", 0.0))
            similarity_percent = round(similarity * 100)
            url = escape(str(src.get("url", "Unknown source")))
            snippet = escape(str(src.get("snippet", "")))
            source_type = src.get("source_type", "internet")
            source_item = (
                f"""
                <div class=\"source-item\">
                    <div><strong>{similarity_percent}% match</strong> - {url}</div>
                    <div class=\"snippet\">{snippet}</div>
                </div>
                """
            )
            if source_type == "internal":
                internal_groups.append(source_item)
            else:
                internet_groups.append(source_item)

    return f"""<!DOCTYPE html>
<html lang=\"en\">
<head>
    <meta charset=\"UTF-8\" />
    <meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\" />
    <title>AI Plagiarism Report</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 24px; color: #111827; }}
        h1 {{ margin-bottom: 6px; }}
        .meta {{ color: #4b5563; margin-bottom: 12px; }}
        .summary {{ display: grid; grid-template-columns: repeat(4, 1fr); gap: 10px; margin-bottom: 20px; }}
        .card {{ border: 1px solid #d1d5db; border-radius: 8px; padding: 10px; background: #f9fafb; }}
        .chart-wrap {{ border: 1px solid #d1d5db; border-radius: 8px; padding: 12px; margin-bottom: 20px; }}
        .bar-row {{ display: flex; align-items: center; margin: 8px 0; gap: 8px; }}
        .bar-label {{ width: 160px; font-weight: 600; }}
        .bar-track {{ flex: 1; background: #e5e7eb; border-radius: 999px; height: 14px; overflow: hidden; }}
        .bar-fill {{ height: 14px; }}
        .bar-fill.plag {{ background: #ef4444; width: {plagiarism_score}%; }}
        .bar-fill.ai {{ background: #10b981; width: {ai_usage_score}%; }}
        .legend {{ margin: 14px 0; font-size: 14px; }}
        .badge {{ display: inline-block; width: 12px; height: 12px; margin-right: 6px; border-radius: 3px; vertical-align: middle; }}
        .b-plag {{ background: rgba(239, 68, 68, 0.35); border: 1px solid #ef4444; }}
        .b-ai {{ background: rgba(16, 185, 129, 0.35); border: 1px solid #10b981; }}
        .b-both {{ background: rgba(245, 158, 11, 0.35); border: 1px solid #f59e0b; }}
        .highlighted {{ line-height: 1.9; border: 1px solid #e5e7eb; border-radius: 8px; padding: 12px; }}
        .sentence {{ padding: 1px 4px; border-radius: 4px; }}
        .sentence.plagiarized {{ background: rgba(239, 68, 68, 0.2); border-bottom: 2px solid #ef4444; }}
        .sentence.ai {{ background: rgba(16, 185, 129, 0.2); border-bottom: 2px solid #10b981; }}
        .sentence.both {{ background: rgba(245, 158, 11, 0.25); border-bottom: 2px solid #f59e0b; }}
        .section {{ margin-top: 24px; }}
        .source-item {{ border: 1px solid #e5e7eb; border-radius: 8px; padding: 10px; margin-bottom: 8px; }}
        .snippet {{ color: #374151; margin-top: 4px; font-style: italic; }}
        table {{ width: 100%; border-collapse: collapse; }}
        th, td {{ border: 1px solid #e5e7eb; padding: 8px; font-size: 14px; text-align: left; vertical-align: top; }}
        th {{ background: #f3f4f6; }}
    </style>
</head>
<body>
    <h1>AI Plagiarism Analysis Report</h1>
    <div class=\"meta\">Generated: {generated_at}</div>

    <div class=\"summary\">
        <div class=\"card\"><strong>Plagiarism Score</strong><br>{plagiarism_score}%</div>
        <div class=\"card\"><strong>AI Usage Score</strong><br>{ai_usage_score}% ({ai_usage_level})</div>
        <div class=\"card\"><strong>Plagiarized Sentences</strong><br>{plagiarized_sentences} / {total_sentences}</div>
        <div class=\"card\"><strong>AI-Likely Sentences</strong><br>{ai_likely_sentences} / {total_sentences}</div>
    </div>

    <div class=\"chart-wrap\">
        <h2>Score Comparison</h2>
        <div class=\"bar-row\"><div class=\"bar-label\">Plagiarism Score</div><div class=\"bar-track\"><div class=\"bar-fill plag\"></div></div><div>{plagiarism_score}%</div></div>
        <div class=\"bar-row\"><div class=\"bar-label\">AI Usage Score</div><div class=\"bar-track\"><div class=\"bar-fill ai\"></div></div><div>{ai_usage_score}%</div></div>
    </div>

    <div class=\"legend\">
        <span class=\"badge b-plag\"></span>Potential Plagiarism
        <span class=\"badge b-ai\" style=\"margin-left: 16px;\"></span>AI-Likely Writing
        <span class=\"badge b-both\" style=\"margin-left: 16px;\"></span>Both Signals
    </div>

    <h2>Highlighted Content</h2>
    <div class=\"highlighted\">{' '.join(rendered_sentences) if rendered_sentences else '<em>No sentence-level result available.</em>'}</div>

    <div class=\"section\">
        <h2>Sentence Explainability</h2>
        <table>
            <thead>
                <tr><th>Sentence</th><th>Plagiarism</th><th>AI-Likely</th><th>Confidence</th><th>Why Flagged</th></tr>
            </thead>
            <tbody>
                {''.join(flagged_details) if flagged_details else '<tr><td colspan="5">No flagged sentences.</td></tr>'}
            </tbody>
        </table>
    </div>

    <div class=\"section\">
        <h2>Internal Database Matches</h2>
        {''.join(internal_groups) if internal_groups else '<p>No internal matches detected.</p>'}
    </div>

    <div class=\"section\">
        <h2>Internet Matches</h2>
        {''.join(internet_groups) if internet_groups else '<p>No internet matches detected.</p>'}
    </div>
</body>
</html>"""

def build_report_pdf(data):
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("PDF export requires reportlab. Install dependencies with requirements.txt.")

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=32, rightMargin=32, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontSize = 9
    title_style = styles["Title"]
    heading = styles["Heading2"]

    # Color palette for visual clarity in PDF exports.
    title_style.textColor = colors.HexColor("#1d4ed8")
    heading.textColor = colors.HexColor("#0f172a")
    heading.spaceBefore = 6
    heading.spaceAfter = 6

    label_style = ParagraphStyle(
        "LabelStyle",
        parent=body,
        textColor=colors.HexColor("#0f172a"),
        fontSize=9,
    )
    label_plag_style = ParagraphStyle(
        "LabelPlagStyle",
        parent=body,
        textColor=colors.HexColor("#b91c1c"),
        fontSize=9,
    )
    label_ai_style = ParagraphStyle(
        "LabelAiStyle",
        parent=body,
        textColor=colors.HexColor("#047857"),
        fontSize=9,
    )
    label_both_style = ParagraphStyle(
        "LabelBothStyle",
        parent=body,
        textColor=colors.HexColor("#b45309"),
        fontSize=9,
    )

    plagiarism_score = round(float(data.get("plagiarism_score", 0) or 0))
    ai_usage_score = round(float(data.get("ai_usage_score", 0) or 0))

    story = [
        Paragraph("AI Plagiarism Analysis Report", title_style),
        Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", body),
        Spacer(1, 10),
    ]

    summary_table = Table([
        ["Plagiarism Score", f"{plagiarism_score}%"],
        ["AI Usage Score", f"{ai_usage_score}%"],
        ["Plagiarized Sentences", f"{int(data.get('plagiarized_sentences', 0))} / {int(data.get('total_sentences', 0))}"],
        ["AI-Likely Sentences", f"{int(data.get('ai_likely_sentences', 0))} / {int(data.get('total_sentences', 0))}"],
    ], colWidths=[170, 330])
    summary_table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#f1f5f9")),
        ("BACKGROUND", (1, 0), (1, 0), colors.HexColor("#fee2e2")),
        ("BACKGROUND", (1, 1), (1, 1), colors.HexColor("#d1fae5")),
        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
        ("TEXTCOLOR", (1, 0), (1, 0), colors.HexColor("#b91c1c")),
        ("TEXTCOLOR", (1, 1), (1, 1), colors.HexColor("#047857")),
    ]))
    story.extend([summary_table, Spacer(1, 12)])

    chart = Drawing(420, 150)
    bar = VerticalBarChart()
    bar.x = 40
    bar.y = 25
    bar.height = 100
    bar.width = 320
    bar.data = [[plagiarism_score, ai_usage_score]]
    bar.categoryAxis.categoryNames = ["Plagiarism", "AI Usage"]
    bar.valueAxis.valueMin = 0
    bar.valueAxis.valueMax = 100
    bar.valueAxis.valueStep = 20
    bar.bars[0].fillColor = colors.HexColor("#ef4444")
    # Color each category bar differently in the same series.
    bar.bars[(0, 0)].fillColor = colors.HexColor("#ef4444")
    bar.bars[(0, 1)].fillColor = colors.HexColor("#10b981")
    bar.valueAxis.strokeColor = colors.HexColor("#94a3b8")
    bar.categoryAxis.strokeColor = colors.HexColor("#94a3b8")
    chart.add(bar)
    story.extend([Paragraph("Score Comparison", heading), chart, Spacer(1, 12)])

    story.append(Paragraph("Highlighted Sentences and Explainability", heading))
    for res in data.get("results", []):
        sentence = escape(str(res.get("sentence", "")))
        is_plagiarized = bool(res.get("is_plagiarized"))
        is_ai_likely = bool(res.get("is_ai_likely"))
        if not (is_plagiarized or is_ai_likely):
            continue

        label = "BOTH" if (is_plagiarized and is_ai_likely) else ("PLAGIARISM" if is_plagiarized else "AI")
        conf = escape(str(res.get("sentence_confidence", "low"))).capitalize()
        reasons = "; ".join([escape(str(r)) for r in res.get("flag_reasons", [])]) or "Heuristic threshold triggered."

        if label == "PLAGIARISM":
            lbl_style = label_plag_style
        elif label == "AI":
            lbl_style = label_ai_style
        elif label == "BOTH":
            lbl_style = label_both_style
        else:
            lbl_style = label_style

        conf_color = "#64748b"
        if conf.lower() == "high":
            conf_color = "#047857"
        elif conf.lower() == "medium":
            conf_color = "#b45309"

        story.append(Paragraph(f"<b>[{label}]</b> {sentence}", lbl_style))
        story.append(Paragraph(f"Confidence: <font color='{conf_color}'><b>{conf}</b></font> | Why: {reasons}", body))
        story.append(Spacer(1, 6))

    story.append(Spacer(1, 12))
    story.append(Paragraph("Internal vs Internet Matches", heading))
    for res in data.get("results", []):
        for src in res.get("sources", []):
            src_type = src.get("source_type", "internet").capitalize()
            sim = round(float(src.get("similarity", 0.0)) * 100)
            url = escape(str(src.get("url", "Unknown source")))
            snippet = escape(str(src.get("snippet", "")))
            story.append(Paragraph(f"{src_type} | {sim}% | {url}", body))
            story.append(Paragraph(f"Snippet: {snippet}", body))
            story.append(Spacer(1, 4))

    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

def sanitize_filename_base(name):
    if not name:
        return "analysis-report"
    clean = re.sub(r"[\\/:*?\"<>|]", "_", str(name)).strip()
    clean = re.sub(r"\s+", " ", clean)
    return clean or "analysis-report"

def get_report_download_name(data, fmt):
    base = sanitize_filename_base(data.get("report_filename_base", "analysis-report"))
    ext = "pdf" if fmt == "pdf" else "html"
    return f"{base}-report.{ext}"

@app.post("/api/report/download")
async def download_report(analysis_json: str = Form(...), format: str = Form("html")):
    try:
        data = json.loads(analysis_json)
    except Exception:
        return {"error": "Invalid report payload."}

    download_name = get_report_download_name(data, format)
    if format == "pdf":
        pdf_bytes = build_report_pdf(data)
        headers = {
            "Content-Disposition": f'attachment; filename="{download_name}"'
        }
        return Response(content=pdf_bytes, media_type="application/pdf", headers=headers)

    report_html = build_report_html(data)
    headers = {
        "Content-Disposition": f'attachment; filename="{download_name}"'
    }
    return Response(content=report_html, media_type="text/html", headers=headers)

@app.get("/api/report/download")
async def download_report_by_token(token: str = Query(...), format: str = Query("html")):
    data = REPORT_CACHE.get(token)
    if not data:
        return {"error": "Report token expired or invalid. Please run analysis again."}

    download_name = get_report_download_name(data, format)
    if format == "pdf":
        pdf_bytes = build_report_pdf(data)
        headers = {
            "Content-Disposition": f'attachment; filename="{download_name}"'
        }
        return Response(content=pdf_bytes, media_type="application/pdf", headers=headers)

    report_html = build_report_html(data)
    headers = {
        "Content-Disposition": f'attachment; filename="{download_name}"'
    }
    return Response(content=report_html, media_type="text/html", headers=headers)

@app.post("/api/check")
async def check_plagiarism(
    text: str = Form(None),
    file: UploadFile = File(None),
    student_name: str = Form(""),
    assignment_name: str = Form(""),
    submission_date: str = Form(""),
):
    try:
        source_text = ""
        report_filename_base = "analysis"
        if file and file.filename:
            content = await file.read()
            report_filename_base = os.path.splitext(file.filename)[0] or "analysis"
            if file.filename.lower().endswith('.pdf'):
                source_text = extract_text_from_pdf(content)
            elif file.filename.lower().endswith('.docx'):
                source_text = extract_text_from_docx(content)
            else:
                try:
                    source_text = content.decode('utf-8')
                except:
                    source_text = ""

            if not source_text.strip():
                return {"error": f"Could not extract any text from '{file.filename}'. It might be scanned, empty, or contain only images."}

        elif text:
            source_text = text

        if not source_text.strip():
            return {"error": "No valid text provided for checking."}

        sentences = split_into_sentences(source_text)
        if not sentences:
            return {"error": "No valid sentences found in the provided text."}

        ensure_model_resources_loaded()

        results = []
        plagiarized_count = 0
        best_similarity_score = 0.0

        user_embeddings = model.encode(sentences)
        semantic_reference_sims = cosine_similarity(user_embeddings, semantic_reference_embeddings)
        semantic_best_similarity = float(np.max(semantic_reference_sims)) if len(sentences) else 0.0

        sentences_to_search = []
        sentence_results = []

        for i, s in enumerate(sentences):
            sentence_result = {
                "sentence": s,
                "is_plagiarized": False,
                "max_similarity": 0.0,
                "sources": []
            }

            user_emb = user_embeddings[i].reshape(1, -1)
            internal_sims = cosine_similarity(user_emb, internal_embeddings)[0]
            max_internal_idx = np.argmax(internal_sims)
            max_internal_sim = float(internal_sims[max_internal_idx])

            # Lower threshold slightly for better recall on paraphrased content
            if max_internal_sim > 0.58:
                sentence_result["is_plagiarized"] = True
                sentence_result["max_similarity"] = max_internal_sim
                sentence_result["sources"].append({
                    "url": "Internal Reference Database",
                    "snippet": INTERNAL_DB[max_internal_idx],
                    "similarity": max_internal_sim,
                    "source_type": "internal",
                })
            else:
                sentences_to_search.append((i, s))

            if sentence_result["max_similarity"] > best_similarity_score:
                best_similarity_score = sentence_result["max_similarity"]

            sentence_results.append(sentence_result)

        if sentences_to_search:
            # Prioritize longer, content-rich chunks for web search to keep latency reasonable.
            if len(sentences_to_search) > MAX_WEB_QUERIES:
                ranked = sorted(
                    sentences_to_search,
                    key=lambda item: len(re.findall(r"\b[a-zA-Z']+\b", item[1])),
                    reverse=True,
                )
                sentences_to_search = ranked[:MAX_WEB_QUERIES]

            async with aiohttp.ClientSession() as session:
                tasks = [search_internet_for_sentence_async(s, session) for _, s in sentences_to_search]
                search_results = await asyncio.gather(*tasks)

                for (idx, s), snippets in zip(sentences_to_search, search_results):
                    if snippets:
                        source_texts = [sn['snippet'] for sn in snippets]
                        source_embs = model.encode(source_texts)
                        user_emb = user_embeddings[idx].reshape(1, -1)
                        web_sims = cosine_similarity(user_emb, source_embs)[0]

                        for j, sim in enumerate(web_sims):
                            if sim > 0.55:
                                sentence_results[idx]["is_plagiarized"] = True
                                if sim > sentence_results[idx]["max_similarity"]:
                                    sentence_results[idx]["max_similarity"] = float(sim)

                                sentence_results[idx]["sources"].append({
                                    "url": snippets[j]['url'],
                                    "snippet": snippets[j]['snippet'],
                                    "similarity": float(sim),
                                    "source_type": "internet",
                                })

                        if sentence_results[idx]["max_similarity"] > best_similarity_score:
                            best_similarity_score = sentence_results[idx]["max_similarity"]

                        sentence_results[idx]["sources"].sort(key=lambda x: x["similarity"], reverse=True)

        plagiarized_count = sum(1 for res in sentence_results if res["is_plagiarized"])
        results = sentence_results
        best_similarity_score = max(best_similarity_score, semantic_best_similarity)

        score = (plagiarized_count / len(sentences)) * 100 if sentences else 0
        ai_usage = analyze_ai_usage(sentences, source_text)

        ai_likely_index_set = set(ai_usage.get("ai_likely_indices", []))
        for idx, sentence_result in enumerate(results):
            sentence_result["is_ai_likely"] = idx in ai_likely_index_set

            ai_details = ai_usage.get("ai_sentence_details", {}).get(idx, {})
            ai_reasons = ai_details.get("reasons", []) if sentence_result["is_ai_likely"] else []
            ai_confidence = ai_details.get("confidence", "low")

            top_similarity = float(sentence_result.get("max_similarity", 0.0))
            plagiarism_reasons = []
            if sentence_result.get("is_plagiarized"):
                plagiarism_reasons.append(f"high semantic similarity ({round(top_similarity * 100)}%)")
                if any(src.get("source_type") == "internal" for src in sentence_result.get("sources", [])):
                    plagiarism_reasons.append("matched internal database")
                if any(src.get("source_type") == "internet" for src in sentence_result.get("sources", [])):
                    plagiarism_reasons.append("matched internet source")

            plagiarism_confidence = confidence_from_similarity(top_similarity)

            level_to_num = {"low": 1, "medium": 2, "high": 3}
            final_confidence = "low"
            if sentence_result.get("is_plagiarized") or sentence_result.get("is_ai_likely"):
                if level_to_num.get(plagiarism_confidence, 1) >= level_to_num.get(ai_confidence, 1):
                    final_confidence = plagiarism_confidence
                else:
                    final_confidence = ai_confidence

            sentence_result["plagiarism_confidence"] = plagiarism_confidence
            sentence_result["ai_confidence"] = ai_confidence
            sentence_result["sentence_confidence"] = final_confidence
            sentence_result["flag_reasons"] = plagiarism_reasons + ai_reasons

        report_token = str(uuid.uuid4())
        report_metadata = {
            "student_name": student_name,
            "assignment_name": assignment_name,
            "submission_date": submission_date,
        }

        response_payload = {
            "plagiarism_score": round(score, 1),
            "best_similarity_score": round(best_similarity_score * 100, 1),
            "semantic_similarity_score": round(semantic_best_similarity * 100, 1),
            "total_sentences": len(sentences),
            "plagiarized_sentences": plagiarized_count,
            "results": results,
            "ai_usage_score": ai_usage["ai_usage_score"],
            "ai_usage_level": ai_usage["ai_usage_level"],
            "ai_likely_sentences": ai_usage["ai_likely_sentences"],
            "ai_likely_indices": ai_usage["ai_likely_indices"],
            "ai_analysis": ai_usage["ai_analysis"],
            "report_metadata": report_metadata,
            "report_filename_base": sanitize_filename_base(report_filename_base),
            "report_token": report_token,
        }

        REPORT_CACHE[report_token] = response_payload
        if len(REPORT_CACHE) > MAX_REPORT_CACHE_ITEMS:
            oldest_token = next(iter(REPORT_CACHE))
            REPORT_CACHE.pop(oldest_token, None)

        return response_payload
    except Exception as e:
        print(f"Error in /api/check: {e}")
        return {"error": f"Analysis failed: {str(e)}"}

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
